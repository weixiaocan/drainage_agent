"""Shared tool runtime: result caching, export/workbook helpers, point and time scope."""

from __future__ import annotations

from copy import deepcopy
import hashlib
import json
import pickle
import re
import time
import traceback
import zipfile
from pathlib import Path
from typing import Any, Callable

import pandas as pd

from agent.deps import AgentDeps
from agent.tools.manifest import data_fingerprint, load_manifest, record_result
from agent.types import FilterConfirmationRequired, ToolResult, error, ok
from analysis import io
from analysis.schema import to_display_columns


SHEET_TABLE_TYPES = {
    "数据体检": "data_check",
    "数据收集率统计": "data_check",
    "日降雨量统计": "rainfall_daily",
    "降雨概况": "rainfall_daily",
    "场次降雨统计": "rainfall_events",
    "降雨场次分析": "rainfall_events",
    "排污规律分析": "patterns",
    "旱天分析": "dry_stats",
    "旱天风险": "dry_risk",
    "雨天溢流风险": "rainy_risk",
    "雨天事件统计": "event_response",
    "RDII总量统计": "rdii",
}


def _rel(deps: AgentDeps, path: Path) -> str:
    try:
        return path.resolve().relative_to(deps.paths.root).as_posix()
    except ValueError:
        return str(path)


def _result_artifacts(deps: AgentDeps, data: dict[str, Any]) -> list[str]:
    """Return only files owned by this tool result, never historical output files."""
    candidates: list[Any] = []
    bundled = False
    for destination in data.get("result_destinations", []):
        if isinstance(destination, dict) and destination.get("path"):
            candidates.append(destination["path"])
            bundled = bundled or destination.get("kind") == "zip"
    if not bundled:
        for key in ("chart_paths", "curve_images", "output_file"):
            if key in data:
                candidates.append(data[key])

    paths: list[str] = []

    def collect(value: Any) -> None:
        if isinstance(value, dict):
            for child in value.values():
                collect(child)
        elif isinstance(value, (list, tuple, set)):
            for child in value:
                collect(child)
        elif isinstance(value, (str, Path)):
            path = Path(value)
            if not path.is_absolute():
                path = deps.paths.root / path
            if path.is_file():
                relative = _rel(deps, path)
                if relative not in paths:
                    paths.append(relative)

    collect(candidates)
    return paths


def _analysis_assets_dir(deps: AgentDeps) -> Path:
    """Keep reusable analysis assets out of the user-facing export directory."""
    path = deps.paths.root / "results" / "generated"
    path.mkdir(parents=True, exist_ok=True)
    return path


_EXPORT_BUNDLE_LABELS = {
    "analyze_patterns": "排污规律分析结果",
    "analyze_rainfall": "降雨分析结果",
    "analyze_rdii": "RDII分析结果",
}


def _bundle_export_artifacts(
    deps: AgentDeps,
    tool_name: str,
    data: dict[str, Any],
    artifacts: list[str],
    params: dict[str, Any],
) -> list[str]:
    """Turn a multi-file explicit export into one stable user download."""
    if not params.get("export") or len(artifacts) <= 1:
        return artifacts
    label = _EXPORT_BUNDLE_LABELS.get(tool_name, "分析结果")
    bundle = deps.paths.outputs / f"{label}.zip"
    bundle.parent.mkdir(parents=True, exist_ok=True)
    internal_root = _analysis_assets_dir(deps)
    with zipfile.ZipFile(bundle, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for artifact in artifacts:
            path = deps.paths.root / artifact
            if not path.is_file() or path == bundle:
                continue
            if path.is_relative_to(deps.paths.outputs):
                arcname = path.relative_to(deps.paths.outputs).as_posix()
            elif path.is_relative_to(internal_root):
                arcname = path.relative_to(internal_root).as_posix()
            else:
                arcname = path.name
            archive.write(path, arcname=arcname)
    for artifact in artifacts:
        path = deps.paths.root / artifact
        if path.is_file() and path.is_relative_to(deps.paths.outputs) and path != bundle:
            path.unlink()
    data["result_destinations"] = [
        {"kind": "zip", "path": _rel(deps, bundle), "sheet": None}
    ]
    return [_rel(deps, bundle)]


def _destination_note(destinations: list[dict[str, Any]]) -> str:
    notes: list[str] = []
    for destination in destinations:
        kind = destination.get("kind")
        path = destination.get("path")
        if kind == "combined_xlsx" and path:
            sheet = destination.get("sheet")
            notes.append(f"已写入 {path}" + (f"（{sheet}）" if sheet else ""))
        elif kind == "csv" and path:
            notes.append(f"已导出 CSV：{path}")
        elif kind == "not_persisted":
            notes.append("本次结果未落盘")
    return "；".join(notes)


class ToolLLMClient:
    def __init__(self, deps: AgentDeps):
        from openai import OpenAI

        kwargs: dict[str, Any] = {"api_key": deps.settings.api_key}
        if deps.settings.base_url:
            kwargs["base_url"] = deps.settings.base_url
        self._client = OpenAI(**kwargs)
        self._model = deps.settings.model
        self._prompt_dirs = [
            deps.paths.root / "agent" / "prompts",
            deps.paths.root / "prompts",
            Path(__file__).resolve().parents[1] / "prompts",
        ]

    def load_prompt(self, name: str) -> str:
        for prompt_dir in self._prompt_dirs:
            path = prompt_dir / f"{name}.txt"
            if path.exists():
                return path.read_text(encoding="utf-8")
        raise FileNotFoundError(name)

    def chat_json(self, prompt: str, temperature: float = 0.1) -> str:
        last_exc: Exception | None = None
        for attempt in range(3):
            try:
                response = self._client.chat.completions.create(
                    model=self._model,
                    messages=[{"role": "user", "content": prompt}],
                    temperature=temperature,
                    response_format={"type": "json_object"},
                )
                return response.choices[0].message.content or "{}"
            except Exception as exc:
                last_exc = exc
                if attempt < 2:
                    time.sleep(2**attempt)
        raise last_exc or RuntimeError("LLM JSON call failed")


def _build_tool_llm_client(deps: AgentDeps) -> ToolLLMClient | None:
    if not deps.settings.api_key:
        return None
    return ToolLLMClient(deps)


def _write_sheet(path: Path, sheet_name: str, df: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    display_df = to_display_columns(df, SHEET_TABLE_TYPES.get(sheet_name, ""))
    mode = "a" if path.exists() else "w"
    if mode == "a":
        with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    else:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    _apply_borders(path, sheet_name)


def _apply_borders(path: Path, sheet_name: str) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = load_workbook(path)
    if sheet_name not in workbook.sheetnames:
        return
    sheet = workbook[sheet_name]
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            cell.border = border
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
    workbook.save(path)


def _site_point_ids(deps: AgentDeps) -> set[str]:
    sites = io.load_sites(root=deps.paths.root)
    for column in ("point_id", "点位编号", "监测点编号", "点位"):
        if column in sites.columns:
            return set(sites[column].dropna().astype(str))
    return set()


def is_full_network(points: list[str] | None, deps: AgentDeps) -> bool:
    if not points:
        return True
    all_points = _site_point_ids(deps)
    if not all_points:
        return False
    values = {str(point).strip() for point in points if str(point).strip()}
    full_scope_aliases = {"全网", "全部点", "全部点位", "所有点", "所有点位"}
    if values.intersection(full_scope_aliases):
        return True
    for value in values:
        match = re.fullmatch(r"(\d+)\s*个?\s*点(?:位)?", value)
        if match and int(match.group(1)) == len(all_points):
            return True
    return all_points.issubset(values)


def _normalize_point_scope(points: list[str] | None, deps: AgentDeps) -> list[str] | None:
    """Use one canonical representation for full-network scope."""
    if is_full_network(points, deps):
        return None
    return list(dict.fromkeys(str(point).strip() for point in points or [] if str(point).strip()))


def is_full_time_range(start: str | None = None, end: str | None = None) -> bool:
    return start is None and end is None


def is_complete_scope(
    points: list[str] | None,
    deps: AgentDeps,
    start: str | None = None,
    end: str | None = None,
) -> bool:
    return is_full_network(points, deps) and is_full_time_range(start, end)


def _safe_filename_part(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", value.strip())
    return cleaned.strip(" ._") or "结果"


def _point_result_prefix(points: list[str] | None) -> str:
    values = sorted({str(point) for point in points or []})
    return "_".join(_safe_filename_part(value) for value in values) or "部分点位"


def _time_result_prefix(start: str | None, end: str | None) -> str:
    if is_full_time_range(start, end):
        return "全时段"

    def format_bound(value: str | None, fallback: str) -> str:
        if value is None:
            return fallback
        parsed = pd.to_datetime(value, errors="coerce")
        if not pd.isna(parsed):
            text = str(value).strip()
            if len(text) > 10:
                return parsed.strftime("%Y-%m-%d_%H-%M-%S")
            return parsed.strftime("%Y-%m-%d")
        return _safe_filename_part(value)

    if start is not None and end is None:
        return f"{format_bound(start, '起始')}_之后"
    if start is None and end is not None:
        return f"{format_bound(end, '结束')}_之前"
    return f"{format_bound(start, '起始')}_{format_bound(end, '结束')}"


def _range_result_prefix(
    points: list[str] | None,
    deps: AgentDeps,
    start: str | None,
    end: str | None,
) -> str:
    point_prefix = "全网" if is_full_network(points, deps) else _point_result_prefix(points)
    return f"{point_prefix}_{_time_result_prefix(start, end)}"


def _route_table_result(
    deps: AgentDeps,
    df: pd.DataFrame,
    sheet_name: str,
    points: list[str] | None,
    export: bool,
    start: str | None = None,
    end: str | None = None,
) -> dict[str, Any]:
    if export:
        filename = f"{_range_result_prefix(points, deps, start, end)}_{_safe_filename_part(sheet_name)}.csv"
        output_path = deps.paths.outputs / filename
        output_path.parent.mkdir(parents=True, exist_ok=True)
        display_df = to_display_columns(df, SHEET_TABLE_TYPES.get(sheet_name, ""))
        display_df.to_csv(output_path, index=False, encoding="utf-8-sig")
        return {"kind": "csv", "path": _rel(deps, output_path), "sheet": None}
    return {"kind": "not_persisted", "path": None, "sheet": None}


REPORT_COMBINED_SHEETS: tuple[tuple[str, str], ...] = (
    ("data_collection", "数据收集率统计"),
    ("rainfall_daily", "降雨概况"),
    ("rainfall_events", "降雨场次分析"),
    ("rainy_event_stats", "雨天事件统计"),
    ("rdii_total", "RDII总量统计"),
    ("pattern_analysis", "排污规律分析"),
    ("dry_analysis", "旱天分析"),
    ("dry_risk", "旱天风险"),
    ("rainy_overflow_risk", "雨天溢流风险"),
)


def _write_report_combined_workbook(
    deps: AgentDeps,
    tables: dict[str, pd.DataFrame],
    output_path: Path,
) -> list[str]:
    """Write only the tables included in the current report."""
    written: list[str] = []
    if output_path.exists():
        output_path.unlink()
    if deps.paths.combined_xlsx != output_path and deps.paths.combined_xlsx.exists():
        deps.paths.combined_xlsx.unlink()
    for key, sheet_name in REPORT_COMBINED_SHEETS:
        table = tables.get(key)
        if table is None or table.empty:
            continue
        _write_sheet(output_path, sheet_name, table)
        written.append(sheet_name)
    return written


def _report_combined_workbook_path(output_file: Path) -> Path:
    report_stem = output_file.stem.removesuffix("_分析报告")
    return output_file.with_name(f"{report_stem}_综合分析结果.xlsx")


def _remove_sheet(path: Path, sheet_name: str) -> None:
    if not path.exists():
        return
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    if sheet_name not in workbook.sheetnames:
        return
    if len(workbook.sheetnames) <= 1:
        path.unlink()
        return
    workbook.remove(workbook[sheet_name])
    workbook.save(path)


def _add_rainfall_excel_charts(path: Path, daily: pd.DataFrame) -> None:
    if daily.empty or not path.exists():
        return
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList

    sheet_name = "降雨概况"
    workbook = load_workbook(path)
    if sheet_name not in workbook.sheetnames:
        return
    sheet = workbook[sheet_name]

    daily_rows = len(daily)
    if daily_rows == 0:
        workbook.save(path)
        return

    pie_col = 5
    rainy_days = int(pd.to_numeric(daily["rain_mm"], errors="coerce").fillna(0).gt(0).sum())
    non_rainy_days = int(daily_rows - rainy_days)
    sheet.cell(row=1, column=pie_col, value="类型")
    sheet.cell(row=1, column=pie_col + 1, value="天数")
    sheet.cell(row=2, column=pie_col, value="降雨日")
    sheet.cell(row=2, column=pie_col + 1, value=rainy_days)
    sheet.cell(row=3, column=pie_col, value="非降雨日")
    sheet.cell(row=3, column=pie_col + 1, value=non_rainy_days)

    bar_chart = BarChart()
    bar_chart.type = "col"
    bar_chart.title = "日降雨量时间序列"
    bar_chart.y_axis.title = "降雨量(mm)"
    bar_chart.x_axis.title = "日期"
    bar_chart.width = 20
    bar_chart.height = 10
    bar_chart.y_axis.majorGridlines = None
    bar_chart.x_axis.majorGridlines = None
    data_ref = Reference(sheet, min_col=2, min_row=1, max_row=daily_rows + 1)
    cats_ref = Reference(sheet, min_col=1, min_row=2, max_row=daily_rows + 1)
    bar_chart.add_data(data_ref, titles_from_data=True)
    bar_chart.set_categories(cats_ref)
    sheet.add_chart(bar_chart, "A8")

    pie_chart = PieChart()
    pie_chart.width = 10
    pie_chart.height = 10
    pie_chart.legend = None
    pie_data = Reference(sheet, min_col=pie_col + 1, min_row=1, max_row=3)
    pie_cats = Reference(sheet, min_col=pie_col, min_row=2, max_row=3)
    pie_chart.add_data(pie_data, titles_from_data=True)
    pie_chart.set_categories(pie_cats)
    pie_chart.dataLabels = DataLabelList()
    pie_chart.dataLabels.showPercent = True
    pie_chart.dataLabels.showVal = True
    pie_chart.dataLabels.showCatName = True
    sheet.add_chart(pie_chart, "A28")

    workbook.save(path)


def _save_pattern_curve_pngs(
    curves: dict[str, pd.DataFrame],
    dry_flow: pd.DataFrame,
    output_dir: Path,
    scope_prefix: str,
) -> dict[str, list[str]]:
    from analysis.pattern_charts import save_pattern_curve_pngs

    return save_pattern_curve_pngs(
        curves,
        dry_flow,
        output_dir,
        scope_prefix,
    )


def _save_partial_pattern_curve_png(
    curves: dict[str, pd.DataFrame],
    dry_flow: pd.DataFrame,
    points: list[str] | None,
    output_dir: Path,
    scope_prefix: str,
) -> dict[str, list[str]]:
    selected_points = {str(value) for value in points or []}
    selected_curves = {
        point_id: curve
        for point_id, curve in curves.items()
        if not selected_points or point_id in selected_points
    }
    return _save_pattern_curve_pngs(
        selected_curves,
        dry_flow,
        output_dir / "特征曲线图",
        scope_prefix,
    )


def _save_partial_rdii_curve_png(
    curve_data: dict[int, dict[str, pd.DataFrame]],
    points: list[str] | None,
    output_dir: Path,
    event_ids: list[int] | None,
) -> dict[int, dict[str, str]]:
    selected_points = {str(point) for point in points or []}
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return {}

    fig = plt.figure(figsize=(10, 5), dpi=120)
    ax = fig.add_subplot(1, 1, 1)
    plotted = False
    for event_id in sorted(curve_data):
        for point_id, frame in sorted(curve_data[event_id].items()):
            if selected_points and str(point_id) not in selected_points:
                continue
            if frame.empty or "rdii_lps" not in frame.columns:
                continue
            values = pd.to_numeric(frame["rdii_lps"], errors="coerce")
            ax.plot(pd.to_datetime(frame.index, errors="coerce"), values, label=f"{point_id}-事件{event_id}")
            plotted = True
    if not plotted:
        plt.close(fig)
        return {}
    output_dir.mkdir(parents=True, exist_ok=True)
    ax.set_xlabel("时间")
    ax.set_ylabel("RDII/(L/s)")
    ax.legend(loc="upper right")
    ax.grid(False)
    fig.tight_layout()
    event_prefix = "_".join(f"event{event_id}" for event_id in sorted(event_ids or [])) or "event未指定"
    path = output_dir / f"{_point_result_prefix(points)}_{event_prefix}_RDII曲线.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return {0: {"selected": str(path)}}


def _run(
    deps: AgentDeps,
    tool_name: str,
    fn: Callable[[], tuple[str, dict[str, Any]]],
    params: dict[str, Any] | None = None,
    use_cache: bool = True,
) -> ToolResult:
    cache_key = _analysis_cache_key(deps, tool_name, params or {})
    if use_cache and cache_key in deps.session.analysis_cache:
        cached = deepcopy(deps.session.analysis_cache[cache_key])
        return ok(cached["summary"], artifacts=_result_artifacts(deps, cached["data"]), **cached["data"])
    try:
        deps.paths.outputs.mkdir(parents=True, exist_ok=True)
        summary, data = fn()
        artifacts = _result_artifacts(deps, data)
        artifacts = _bundle_export_artifacts(
            deps, tool_name, data, artifacts, params or {}
        )
        destination_note = _destination_note(data.get("result_destinations", []))
        if destination_note:
            summary = f"{summary} 落盘去向：{destination_note}。"
        record_result(deps, tool_name, artifacts, params=params)
        if use_cache:
            deps.session.analysis_cache[cache_key] = {"summary": summary, "data": deepcopy(data)}
        _store_report_data_components(deps, tool_name, params or {}, data)
        return ok(summary, artifacts=artifacts, **data)
    except FilterConfirmationRequired:
        raise
    except Exception as exc:
        deps.logger.exception("%s failed", tool_name)
        return error(f"{tool_name} 执行失败: {exc}", traceback=traceback.format_exc(limit=8))


def _analysis_cache_key(deps: AgentDeps, tool_name: str, params: dict[str, Any]) -> str:
    normalized_params = dict(params)
    for key in ("points", "event_ids", "sections"):
        value = normalized_params.get(key)
        if isinstance(value, list):
            normalized_params[key] = sorted(value, key=str)
    payload = {
        "tool": tool_name,
        "params": normalized_params,
        "data_fingerprint": data_fingerprint(deps)["digest"],
    }
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)


def _report_data_key(
    deps: AgentDeps,
    component: str,
    points: list[str] | None = None,
    start: str | None = None,
    end: str | None = None,
    event_ids: list[int] | None = None,
) -> str:
    payload = {
        "component": component,
        "points": sorted(points or [], key=str),
        "start": start,
        "end": end,
        "event_ids": sorted(event_ids or []),
        "data_fingerprint": data_fingerprint(deps)["digest"],
    }
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)


def _store_report_data_components(
    deps: AgentDeps,
    tool_name: str,
    params: dict[str, Any],
    data: dict[str, Any],
) -> None:
    points = params.get("points") or []
    start = params.get("start")
    end = params.get("end")
    event_ids = params.get("event_ids") or []
    components: list[tuple[str, str, list[int] | None]] = []
    if tool_name == "check_data":
        components = [("data_collection", "table", None)]
    elif tool_name == "analyze_rainfall":
        time_range = params.get("time_range") or []
        start = time_range[0] if len(time_range) == 2 else None
        end = time_range[1] if len(time_range) == 2 else None
        components = [
            ("rainfall_daily", "daily", None),
            ("rainfall_events", "events", None),
            ("rainfall_chart_paths", "chart_paths", None),
        ]
    elif tool_name == "analyze_patterns":
        components = [
            ("pattern_analysis", "table", None),
            ("pattern_chart_paths", "curve_images", None),
        ]
    elif tool_name == "assess_risk":
        scope = params.get("scope", "all")
        if scope in {"dry", "all"} and data.get("dry_analysis") is not None:
            components.append(("dry_analysis", "dry_analysis", None))
        if scope in {"dry", "all"} and data.get("dry_risk") is not None:
            components.append(("dry_risk", "dry_risk", None))
        if scope in {"rainy", "all"} and data.get("rainy_risk") is not None:
            components.append(("rainy_overflow_risk", "rainy_risk", event_ids))
    for component, data_key, component_events in components:
        records = data.get(data_key)
        if records is None:
            continue
        key = _report_data_key(deps, component, points, start, end, component_events)
        deps.session.report_data_cache[key] = deepcopy(records)


def _cached_report_frame(
    deps: AgentDeps,
    component: str,
    points: list[str] | None = None,
    start: str | None = None,
    end: str | None = None,
    event_ids: list[int] | None = None,
) -> pd.DataFrame | None:
    key = _report_data_key(deps, component, points, start, end, event_ids)
    records = deps.session.report_data_cache.get(key)
    return pd.DataFrame(deepcopy(records)) if records is not None else None


def _cached_report_value(
    deps: AgentDeps,
    component: str,
    points: list[str] | None = None,
    start: str | None = None,
    end: str | None = None,
    event_ids: list[int] | None = None,
) -> Any | None:
    key = _report_data_key(deps, component, points, start, end, event_ids)
    value = deps.session.report_data_cache.get(key)
    return deepcopy(value) if value is not None else None


def _manifest_stale(deps: AgentDeps, tool_name: str) -> bool:
    manifest = load_manifest(deps)
    item = manifest.get("results", {}).get(tool_name)
    if not item:
        return True
    from agent.tools.manifest import data_fingerprint

    return item.get("data_fingerprint") != data_fingerprint(deps)["digest"]


def _file_sha256(path: Path) -> str | None:
    if not path.exists() or not path.is_file():
        return None
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()



def _intermediate_dir(deps: AgentDeps) -> Path:
    path = _analysis_assets_dir(deps) / "intermediate"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _curves_path(deps: AgentDeps) -> Path:
    return _intermediate_dir(deps) / "dry_curves.pkl"


def _rdii_curves_path(deps: AgentDeps) -> Path:
    return _intermediate_dir(deps) / "rdii_curves.pkl"


def _save_curves(deps: AgentDeps, curves: dict[str, pd.DataFrame]) -> None:
    with _curves_path(deps).open("wb") as fh:
        pickle.dump(curves, fh)


def _save_rdii_curves(deps: AgentDeps, curves: dict[int, dict[str, pd.DataFrame]]) -> None:
    with _rdii_curves_path(deps).open("wb") as fh:
        pickle.dump(curves, fh)


def _load_curves(deps: AgentDeps) -> dict[str, pd.DataFrame]:
    path = _curves_path(deps)
    if not path.exists():
        return {}
    with path.open("rb") as fh:
        return pickle.load(fh)
